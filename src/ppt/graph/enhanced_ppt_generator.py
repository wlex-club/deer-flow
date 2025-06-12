# Copyright (c) 2025 Bytedance Ltd. and/or its affiliates
# SPDX-License-Identifier: MIT

import logging
import os
import subprocess
import uuid
import json
import re
from pathlib import Path
from typing import Optional
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

from src.ppt.graph.state import PPTState
from src.utils.file_manager import file_manager

logger = logging.getLogger(__name__)


def extract_title_from_content(content: str) -> str:
    """从内容中提取标题"""
    lines = content.strip().split('\n')
    for line in lines:
        line = line.strip()
        if line.startswith('# '):
            # 移除markdown标记和emoji
            title = re.sub(r'^#\s*', '', line)
            title = re.sub(r'[^\w\s-]', '', title).strip()
            return title[:50]  # 限制长度
    return "presentation"


def parse_markdown_to_slides(markdown_content: str) -> list:
    """解析Markdown内容为幻灯片数据"""
    slides = []
    current_slide = None
    
    lines = markdown_content.strip().split('\n')
    
    for line in lines:
        line = line.strip()
        
        # 跳过YAML frontmatter
        if line.startswith('---') and len(slides) == 0:
            continue
        
        # 新幻灯片分隔符
        if line == '---' and current_slide is not None:
            if current_slide['title'] or current_slide['content']:
                slides.append(current_slide)
            current_slide = {'title': '', 'content': [], 'level': 1}
            continue
        
        # 标题
        if line.startswith('# '):
            if current_slide is not None:
                slides.append(current_slide)
            current_slide = {
                'title': line[2:].strip(),
                'content': [],
                'level': 1
            }
        elif line.startswith('## '):
            if current_slide is not None:
                slides.append(current_slide)
            current_slide = {
                'title': line[3:].strip(),
                'content': [],
                'level': 2
            }
        elif line.startswith('### '):
            if current_slide is not None:
                slides.append(current_slide)
            current_slide = {
                'title': line[4:].strip(),
                'content': [],
                'level': 3
            }
        # 内容
        elif line and current_slide is not None:
            current_slide['content'].append(line)
        elif current_slide is None:
            # 第一张幻灯片
            current_slide = {
                'title': line if line else 'Presentation',
                'content': [],
                'level': 1
            }
    
    # 添加最后一张幻灯片
    if current_slide is not None and (current_slide['title'] or current_slide['content']):
        slides.append(current_slide)
    
    return slides


def create_powerpoint_presentation(slides_data: list, title: str) -> str:
    """使用python-pptx创建PowerPoint演示文稿"""
    
    # 创建演示文稿
    prs = Presentation()
    
    # 设置幻灯片大小 (16:9)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    for i, slide_data in enumerate(slides_data):
        # 选择布局
        if i == 0:
            # 标题幻灯片
            slide_layout = prs.slide_layouts[0]  # Title Slide
        else:
            # 内容幻灯片
            slide_layout = prs.slide_layouts[1]  # Title and Content
        
        slide = prs.slides.add_slide(slide_layout)
        
        # 设置标题
        if slide.shapes.title:
            slide.shapes.title.text = slide_data['title']
            
            # 设置标题样式
            title_shape = slide.shapes.title
            title_frame = title_shape.text_frame
            title_frame.clear()
            
            p = title_frame.paragraphs[0]
            p.text = slide_data['title']
            p.alignment = PP_ALIGN.CENTER if i == 0 else PP_ALIGN.LEFT
            
            # 设置字体
            font = p.font
            font.name = 'Microsoft YaHei'
            font.size = Pt(36) if i == 0 else Pt(28)
            font.bold = True
            font.color.rgb = RGBColor(44, 62, 80)  # 深蓝色
        
        # 设置内容
        if len(slide_data['content']) > 0 and len(slide.placeholders) > 1:
            content_placeholder = slide.placeholders[1]
            text_frame = content_placeholder.text_frame
            text_frame.clear()
            
            for j, content_line in enumerate(slide_data['content']):
                if j == 0:
                    p = text_frame.paragraphs[0]
                else:
                    p = text_frame.add_paragraph()
                
                # 处理列表项
                if content_line.startswith('- ') or content_line.startswith('* '):
                    p.text = content_line[2:]
                    p.level = 0
                elif content_line.startswith('  - ') or content_line.startswith('  * '):
                    p.text = content_line[4:]
                    p.level = 1
                elif content_line.startswith(('1. ', '2. ', '3. ', '4. ', '5. ')):
                    p.text = content_line[3:]
                    p.level = 0
                else:
                    p.text = content_line
                    p.level = 0
                
                # 设置字体
                font = p.font
                font.name = 'Microsoft YaHei'
                font.size = Pt(18)
                font.color.rgb = RGBColor(52, 73, 94)  # 深灰色
                
                # 处理加粗和斜体
                if '**' in content_line:
                    # 简单的加粗处理
                    font.bold = True
    
    # 保存文件
    filename = file_manager.generate_filename("presentation", title, "pptx")
    output_path = file_manager.get_output_path("pptx", filename)
    
    try:
        prs.save(output_path)
        logger.info(f"PowerPoint presentation saved: {output_path}")
        return str(output_path)
    except Exception as e:
        logger.error(f"Failed to save PowerPoint presentation: {e}")
        raise


def generate_html_presentation_enhanced(slides_data: list, title: str) -> str:
    """生成增强的HTML演示文稿"""
    
    # 生成文件名
    filename = file_manager.generate_filename("presentation", title, "html")
    output_path = file_manager.get_output_path("html", filename)
    
    # HTML模板
    html_template = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{title}</title>
    <link rel="stylesheet" href="https://cdn.jsdelivr.net/npm/reveal.js@4.3.1/dist/reveal.css">
    <link rel="stylesheet" href="https://cdn.jsdelivr.net/npm/reveal.js@4.3.1/dist/theme/white.css">
    <style>
        .reveal {{
            font-family: 'Microsoft YaHei', 'SimHei', Arial, sans-serif;
        }}
        
        .reveal .slides section {{
            text-align: left;
            background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
            color: white;
            padding: 40px;
            border-radius: 15px;
            box-shadow: 0 10px 30px rgba(0,0,0,0.3);
        }}
        
        .reveal .slides section.title-slide {{
            text-align: center;
            background: linear-gradient(135deg, #4facfe 0%, #00f2fe 100%);
        }}
        
        .reveal h1 {{
            font-size: 2.5em;
            font-weight: bold;
            margin-bottom: 30px;
            text-shadow: 2px 2px 4px rgba(0,0,0,0.3);
        }}
        
        .reveal h2 {{
            font-size: 2em;
            font-weight: bold;
            margin-bottom: 25px;
            color: #ffd700;
        }}
        
        .reveal ul {{
            list-style: none;
            padding-left: 0;
        }}
        
        .reveal ul li {{
            margin: 15px 0;
            padding-left: 30px;
            position: relative;
            line-height: 1.6;
        }}
        
        .reveal ul li::before {{
            content: "●";
            position: absolute;
            left: 0;
            color: #ffd700;
            font-size: 1.2em;
        }}
        
        .reveal strong {{
            color: #ffd700;
            font-weight: bold;
        }}
        
        .navigation-hint {{
            position: fixed;
            bottom: 20px;
            right: 20px;
            background: rgba(0,0,0,0.7);
            color: white;
            padding: 10px 15px;
            border-radius: 5px;
            font-size: 14px;
            z-index: 1000;
        }}
    </style>
</head>
<body>
    <div class="reveal">
        <div class="slides">
"""
    
    # 生成幻灯片内容
    slides_html = ""
    for i, slide_data in enumerate(slides_data):
        slide_class = "title-slide" if i == 0 else ""
        slides_html += f'<section class="{slide_class}">\n'
        
        if i == 0:
            slides_html += f'<h1>{slide_data["title"]}</h1>\n'
        else:
            slides_html += f'<h2>{slide_data["title"]}</h2>\n'
        
        if slide_data['content']:
            slides_html += '<ul>\n'
            for content_line in slide_data['content']:
                # 处理markdown格式
                content_line = content_line.replace('**', '<strong>').replace('**', '</strong>')
                content_line = content_line.replace('*', '<em>').replace('*', '</em>')
                
                if content_line.startswith(('- ', '* ', '1. ', '2. ', '3. ')):
                    clean_line = re.sub(r'^[-*\d.]\s*', '', content_line)
                    slides_html += f'<li>{clean_line}</li>\n'
                else:
                    slides_html += f'<li>{content_line}</li>\n'
            slides_html += '</ul>\n'
        
        slides_html += '</section>\n'
    
    # 完整HTML
    full_html = html_template + slides_html + """
        </div>
    </div>
    
    <div class="navigation-hint">
        使用方向键或空格键切换幻灯片
    </div>
    
    <script src="https://cdn.jsdelivr.net/npm/reveal.js@4.3.1/dist/reveal.js"></script>
    <script>
        Reveal.initialize({
            hash: true,
            controls: true,
            progress: true,
            center: false,
            transition: 'slide',
            transitionSpeed: 'default'
        });
    </script>
</body>
</html>"""
    
    # 写入文件
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(full_html)
    
    logger.info(f"HTML presentation generated: {output_path}")
    return str(output_path)


def enhanced_ppt_generator_node(state: PPTState):
    """增强的PPT生成器节点"""
    logger.info("Generating enhanced presentation...")
    
    try:
        # 读取markdown内容
        with open(state["ppt_file_path"], 'r', encoding='utf-8') as f:
            markdown_content = f.read()
        
        # 提取标题
        title = extract_title_from_content(markdown_content)
        
        # 解析幻灯片数据
        slides_data = parse_markdown_to_slides(markdown_content)
        
        if not slides_data:
            logger.warning("No slides found in markdown content")
            slides_data = [{'title': title, 'content': ['No content found'], 'level': 1}]
        
        logger.info(f"Parsed {len(slides_data)} slides")
        
        # 尝试生成PowerPoint文件
        try:
            pptx_path = create_powerpoint_presentation(slides_data, title)
            
            # 验证文件
            if os.path.exists(pptx_path) and os.path.getsize(pptx_path) > 1000:
                logger.info(f"Successfully generated PPTX: {pptx_path}")
                
                # 清理临时文件
                if os.path.exists(state["ppt_file_path"]):
                    os.remove(state["ppt_file_path"])
                
                return {"generated_file_path": pptx_path}
            else:
                logger.warning("PPTX file seems corrupted, falling back to HTML")
                raise Exception("PPTX file validation failed")
                
        except Exception as e:
            logger.warning(f"PowerPoint generation failed: {e}, falling back to HTML")
            
            # 生成HTML作为备选方案
            html_path = generate_html_presentation_enhanced(slides_data, title)
            
            # 清理临时文件
            if os.path.exists(state["ppt_file_path"]):
                os.remove(state["ppt_file_path"])
            
            return {"generated_file_path": html_path}
            
    except Exception as e:
        logger.error(f"Enhanced PPT generation failed: {e}")
        raise Exception(f"Failed to generate presentation: {e}") 